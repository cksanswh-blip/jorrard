#!/usr/bin/env python3
"""OrangPro 로드맵 — git 정본 도구.

2026-09-24 부터 **정본은 이 저장소의 TSV** 다(대표 결정). 시트 `Orang정보총망라`·`총책임자전용` 은
사람이 보는 거울이고, 거울은 git 에서 붙여넣기(또는 브릿지)로 맞춘다. 시트에서 먼저 고치지 않는다.

  add TAB DELTA.tsv      새 행을 TSV 에 붙이고 검사한 뒤 행 날짜로 커밋한다 — **평소 기입은 이것 하나**
  check                  TSV 3장을 검사한다(열 수 · 번호 중복·순서 · 날짜 형식). 커밋 전에 돈다
  summary                TSV 에서 SUMMARY.md 를 다시 쓴다(커밋 안 함)

  (시트에서 가져오기 — 2026-09-24 이전 이력과 거울 대조용)
  export   XLSX            작업트리에 TSV·SUMMARY.md 를 쓴다(커밋 안 함)

  export   XLSX            작업트리에 TSV·SUMMARY.md 를 쓴다(커밋 안 함)
  backfill XLSX           작성·수정일 하루 1커밋의 소급 사슬을 부모 없이 만들고 끝 SHA 를 출력한다.
                          현재 브랜치에는 `git merge --allow-unrelated-histories <SHA>` 로 붙인다
  sync     XLSX           HEAD 의 TSV 와 비교해 바뀐 행만 그 행의 날짜로 소급 커밋한다(시트가 앞서갔을 때만)

XLSX 는 구글 시트 「파일 → 다운로드 → Microsoft Excel」 로 받은 파일이다.
공개 저장소를 전제로 기본 마스킹이 켜져 있다. 비공개 저장소에서만 --no-redact 를 쓴다.
"""
import argparse
import collections
import datetime as dt
import os
import re
import subprocess
import sys
import tempfile

OUT_DIR = "orangpro/roadmap"
KST = "+0900"

# 탭 이름 → (저장 파일명, 날짜 열 이름)
# 날짜 열이 없으면 None: 커밋 시각(또는 add --date)을 쓴다
TABS = {
    "로드맵": ("로드맵.tsv", "작성·수정일"),
    "로드맵(미래)": ("로드맵_미래.tsv", "작성·수정일"),
    "100억": ("총책임자전용_100억.tsv", None),   # 시트 `총책임자전용` 의 탭
}

AUTHOR = ("Orang정보총망라 (시트 소급)", "roadmap@orangpro.invalid")   # 시트에서 가져온 행
AUTHOR_GIT = ("OrangPro 로드맵 (git 정본)", "roadmap@orangpro.invalid")  # git 에 바로 적은 행
COMMITTER = ("Claude", "noreply@anthropic.com")
TRAILER = (
    "\n\nCo-Authored-By: Claude Fable 5.1 <noreply@anthropic.com>\n"
    "Claude-Session: https://claude.ai/code/session_01F3Wvgo5RLKV6FLU27p5JZv\n"
)

# ── 마스킹 ───────────────────────────────────────────────────────────────
# 공개 저장소에 나가면 안 되는 것만 가린다. 해시(sha16)·WP 페이지 번호·파일명은 남긴다.
REDACT = [
    (re.compile(r"AKfyc[\w-]+"), "[GAS배포ID]"),
    (re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+"), "[이메일]"),
    (re.compile(r"\b19\d\d-\d\d-\d\d\b"), "[생년월일]"),
    (re.compile(r"\b4-\d{4}-\d{6}-\d\b"), "[특허고객번호]"),
    (re.compile(r"\b\d{3}-\d{2}-\d{5}\b"), "[사업자등록번호]"),
    # 드라이브·시트 ID: 영문·숫자가 섞인 28자 이상의 연속 토큰
    (re.compile(r"(?<![\w/.-])(?=[\w-]*\d)(?=[\w-]*[A-Za-z])[\w-]{28,}(?![\w.])"), "[드라이브ID]"),
]

# 실명처럼 코드에 적으면 그 자체로 노출되는 규칙은 커밋하지 않는 로컬 파일에 둔다.
# 한 줄에 「정규식<TAB>바꿀말」. 파일은 .gitignore 에 걸려 있다.
LOCAL_RULES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "redact_local.tsv")
if os.path.exists(LOCAL_RULES):
    with open(LOCAL_RULES, encoding="utf-8") as _f:
        for _ln in _f:
            if _ln.strip() and not _ln.startswith("#") and "\t" in _ln:
                _p, _r = _ln.rstrip("\n").split("\t", 1)
                REDACT.append((re.compile(_p), _r))


def redact(s):
    for pat, rep in REDACT:
        s = pat.sub(rep, s)
    return s


# ── 읽기 ─────────────────────────────────────────────────────────────────
def cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    s = str(v).strip()
    return s.replace("\r\n", "\n").replace("\n", " ⏎ ").replace("\t", "    ")


def parse_when(s):
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s[:19].strip(), fmt)
        except ValueError:
            pass
    return None


def load(xlsx, do_redact):
    """탭별 (header, rows) 를 돌려준다. rows 는 번호 순 (key, cells, when)."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl 이 필요하다: pip install openpyxl")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out = {}
    for tab, (_, date_col) in TABS.items():
        if tab not in wb.sheetnames:
            continue
        it = wb[tab].iter_rows(values_only=True)
        header = [cell(c) for c in next(it)]
        while header and not header[-1]:
            header.pop()
        width = len(header)
        di = header.index(date_col) if date_col else None
        rows, last_when = [], None
        for raw in it:
            cells = [cell(c) for c in list(raw)[:width]]
            cells += [""] * (width - len(cells))
            if not any(cells):
                continue
            if do_redact:
                cells = [redact(c) for c in cells]
            when = parse_when(cells[di]) if di is not None and cells[di] else None
            # 날짜가 빈 행은 시트에서 바로 앞 행의 날짜를 물려받는다(번호가 곧 기입 순서라는 가정)
            if when is None:
                when = last_when
            else:
                last_when = when
            rows.append([cells[0], cells, when])
        first = next((r[2] for r in rows if r[2]), dt.datetime(2026, 7, 1))
        for r in rows:
            r[2] = r[2] or first
        rows.sort(key=lambda r: (int(r[0]) if r[0].isdigit() else 10**9, r[0]))
        out[tab] = (header, rows)
    return out


def render(header, rows):
    lines = ["\t".join(header)] + ["\t".join(r[1]) for r in rows]
    return "\n".join(lines) + "\n"


def read_tsv(path):
    if not os.path.exists(path):
        return None, {}
    with open(path, encoding="utf-8") as f:
        lines = f.read().rstrip("\n").split("\n")
    header = lines[0].split("\t")
    rows = {}
    for ln in lines[1:]:
        cells = ln.split("\t")
        rows[cells[0]] = cells
    return header, rows


def row_when(tab, header, cells, fallback):
    date_col = TABS[tab][1]
    if date_col and date_col in header:
        v = cells[header.index(date_col)]
        w = parse_when(v) if v else None
        if w:
            return w
    return fallback


def sort_key(k):
    return (int(k) if k.isdigit() else 10**9, k)


def load_tsv_all():
    """git 작업트리의 TSV 를 load() 와 같은 모양 {탭: (header, [key, cells, when])} 으로 돌려준다."""
    out = {}
    for tab, (fname, _) in TABS.items():
        header, rows = read_tsv(os.path.join(OUT_DIR, fname))
        if header is None:
            continue
        lst, last = [], dt.datetime(2026, 7, 1)
        for k in sorted(rows, key=sort_key):
            when = row_when(tab, header, rows[k], last)
            last = when
            lst.append([k, rows[k], when])
        out[tab] = (header, lst)
    return out


def check_tab(tab, header, rows, strict=True):
    """오류 문장 목록을 돌려준다. rows 는 [key, cells, when]."""
    errs, width, seen, prev = [], len(header), set(), 0
    date_col = TABS[tab][1]
    for k, cells, _ in rows:
        if len(cells) != width:
            errs.append(f"{tab} {k}번: 열 {len(cells)}개 (헤더 {width}개)")
        if not k.isdigit():
            errs.append(f"{tab} {k!r}: 번호가 숫자가 아니다")
            continue
        if k in seen:
            errs.append(f"{tab} {k}번: 번호 중복")
        seen.add(k)
        if strict and int(k) != prev + 1:
            errs.append(f"{tab} {k}번: 번호가 이어지지 않는다 (앞 {prev})")
        prev = int(k)
        if date_col and date_col in header:
            v = cells[header.index(date_col)]
            if v and not re.match(r"\d{4}-\d{2}-\d{2}", v):
                errs.append(f"{tab} {k}번: 날짜가 YYYY-MM-DD 로 시작하지 않는다 {v!r}")
        if any(("\n" in c or "\t" in c) for c in cells):
            errs.append(f"{tab} {k}번: 셀 안 줄바꿈·탭 — ' ⏎ ' 와 공백 4칸으로 바꾼다")
    return errs


def cmd_check(a):
    data = load_tsv_all()
    if not data:
        sys.exit(f"{OUT_DIR} 에 TSV 가 없다")
    errs = []
    for tab, (header, rows) in data.items():
        errs += check_tab(tab, header, rows, strict=not a.loose)
        print(f"{tab}: {len(rows)}행 · 마지막 {rows[-1][0] if rows else '-'}번")
    if errs:
        sys.exit("\n".join("✗ " + e for e in errs))
    print("✓ 이상 없음")


def cmd_add(a):
    """DELTA.tsv(헤더 없음, 열 수 동일) 의 행을 TAB 에 붙이고 검사한 뒤 행 날짜별로 커밋한다."""
    tab = a.tab
    if tab not in TABS:
        sys.exit(f"탭은 {', '.join(TABS)} 중 하나다")
    if git("status", "--porcelain", "--", OUT_DIR):
        sys.exit(f"{OUT_DIR} 에 커밋 안 된 변경이 있다. 먼저 정리한다.")
    path = os.path.join(OUT_DIR, TABS[tab][0])
    header, old = read_tsv(path)
    if header is None:
        sys.exit(f"{path} 가 없다. 첫 적재는 export 로 만든다")
    width = len(header)
    with open(a.delta, encoding="utf-8") as f:
        lines = [ln for ln in f.read().rstrip("\n").split("\n") if ln.strip()]
    if lines and lines[0].split("\t")[0] == header[0]:
        lines = lines[1:]  # 헤더가 들어 있으면 버린다
    fallback = dt.datetime.strptime(a.date, "%Y-%m-%d") if a.date else dt.datetime.now()
    new_rows = []
    for ln in lines:
        cells = [cell(c) for c in ln.split("\t")]
        if len(cells) > width:
            sys.exit(f"{cells[0]}번: 열 {len(cells)}개 — 헤더는 {width}개")
        cells += [""] * (width - len(cells))
        if not a.no_redact:
            cells = [redact(c) for c in cells]
        if cells[0] in old:
            sys.exit(f"{cells[0]}번은 이미 있다. 고치려면 TSV 를 직접 고치고 커밋한다")
        new_rows.append([cells[0], cells, row_when(tab, header, cells, fallback)])
    if not new_rows:
        sys.exit("붙일 행이 없다")
    merged = dict(old)
    merged.update({k: c for k, c, _ in new_rows})
    rows_all, last = [], dt.datetime(2026, 7, 1)
    for k in sorted(merged, key=sort_key):
        w = row_when(tab, header, merged[k], last)
        last = w
        rows_all.append([k, merged[k], w])
    errs = check_tab(tab, header, rows_all, strict=not a.loose)
    if errs:
        sys.exit("\n".join("✗ " + e for e in errs) + "\n아무것도 쓰지 않았다")
    print(f"{tab}: +{len(new_rows)}행 ({new_rows[0][0]}~{new_rows[-1][0]}번) → {len(rows_all)}행")
    if a.dry_run:
        return
    by_day = collections.defaultdict(list)
    for r in new_rows:
        by_day[r[2].date()].append(r)
    written = dict(old)
    for day in sorted(by_day):
        for r in by_day[day]:
            written[r[0]] = r[1]
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(["\t".join(header)] + ["\t".join(written[k]) for k in sorted(written, key=sort_key)]) + "\n")
        nums = [r[0] for r in by_day[day]]
        latest = max(r[2] for r in by_day[day])
        title = f"로드맵 기입 {day}: {tab} +{len(nums)} ({nums[0]}~{nums[-1]}번)"
        body = "\n\n" + "\n".join(f"- {r[0]} {r[1][2][:70]}" for r in by_day[day])
        if a.why:
            body += f"\n\n{a.why}"
        env = dated_env(latest)
        env["GIT_AUTHOR_NAME"], env["GIT_AUTHOR_EMAIL"] = AUTHOR_GIT
        git("add", path)
        git("commit", "-q", "-m", title + body + TRAILER, env=env)
        print(f"  {day}  +{len(nums)}  커밋")
    write_summary(load_tsv_all())
    if git("status", "--porcelain", "--", OUT_DIR):
        git("add", OUT_DIR)
        git("commit", "-q", "-m", "로드맵 요약 갱신" + TRAILER)
    print(f"완료 — 거울(시트)에는 같은 행을 붙여넣는다: {a.delta}")


def cmd_summary(a):
    write_summary(load_tsv_all())
    print("SUMMARY.md 갱신")


def write_summary(data):
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "SUMMARY.md"), "w", encoding="utf-8") as f:
        f.write(summary(data))

# ── 요약 ─────────────────────────────────────────────────────────────────
def status_bucket(s):
    if s.startswith("✅") or s.startswith("완료") or "완료(" in s[:8]:
        return "✅ 완료"
    if "진행" in s or s.startswith("🔧") or s.startswith("부분") or "운영중" in s:
        return "🔧 진행·운영"
    if s.startswith("⏳") or "준비" in s or "대기" in s or s.startswith("계획"):
        return "⏳ 준비·대기·계획"
    if "사고" in s or "⚠" in s or "점검" in s:
        return "⚠️ 점검·사고"
    return "· 기타"


def summary(data):
    L = ["# 로드맵 요약 (자동 생성 — 손으로 고치지 않는다)", ""]
    L.append("`python3 tools/roadmap_git.py add|summary|sync` 가 매번 다시 쓴다. **정본은 이 저장소의 TSV** (2026-09-24부터). 시트는 거울.")
    L.append("")
    if "로드맵" in data:
        header, rows = data["로드맵"]
        L += [f"## 로드맵 — {len(rows)}행 (구현·완료 이력)", ""]
        months = collections.Counter(r[2].strftime("%Y-%m") for r in rows)
        L += ["| 월 | 행 |", "| :-- | --: |"] + [f"| {m} | {n} |" for m, n in sorted(months.items())]
        L.append("")
        st = collections.Counter(status_bucket(r[1][4]) for r in rows)
        L += ["| 상태 | 행 |", "| :-- | --: |"] + [f"| {k} | {v} |" for k, v in st.most_common()]
        L.append("")
        area = collections.Counter(r[1][1] for r in rows)
        L += ["| 영역 (상위 25) | 행 |", "| :-- | --: |"] + [f"| {k} | {v} |" for k, v in area.most_common(25)]
        L.append("")
        L += ["### 최근 20행", "", "| 번호 | 영역 | 컴포넌트·기능 | 상태 | 날짜 |", "| --: | :-- | :-- | :-- | :-- |"]
        for r in sorted(rows, key=lambda r: r[2])[-20:][::-1]:
            c = r[1]
            L.append(f"| {c[0]} | {c[1]} | {c[2][:60]} | {c[4][:20]} | {r[2]:%Y-%m-%d} |")
        L.append("")
    if "로드맵(미래)" in data:
        header, rows = data["로드맵(미래)"]
        L += [f"## 로드맵(미래) — {len(rows)}행 (계획·기둥)", ""]
        pr = collections.Counter(r[1][5] or "·" for r in rows)
        L += ["| 우선순위 | 행 |", "| :-- | --: |"] + [f"| {k} | {v} |" for k, v in sorted(pr.items())]
        L.append("")
        L += ["### 열린 P0", "", "| 번호 | 기둥 | 핵심 기능 | 상태 |", "| --: | :-- | :-- | :-- |"]
        for r in rows:
            c = r[1]
            if c[5] == "P0" and status_bucket(c[6]) != "✅ 완료":
                L.append(f"| {c[0]} | {c[1]} | {c[2][:50]} | {c[6][:40]} |")
        L.append("")
    if "100억" in data:
        header, rows = data["100억"]
        L += [f"## 총책임자전용 · 100억 — {len(rows)}행 (자격·트랙션·정부지원·IP·행정)", ""]
        tr = collections.Counter(r[1][1] for r in rows)
        L += ["| 트랙 | 행 |", "| :-- | --: |"] + [f"| {k} | {v} |" for k, v in sorted(tr.items())]
        L.append("")
        st = collections.Counter(status_bucket(r[1][5]) for r in rows)
        L += ["| 상태 | 행 |", "| :-- | --: |"] + [f"| {k} | {v} |" for k, v in st.most_common()]
        L.append("")
        L += ["### 열린 자격·정부지원 행 (대기·준비)", "", "| 번호 | 트랙 | 항목 | 시점·마감 |", "| --: | :-- | :-- | :-- |"]
        for r in rows:
            c = r[1]
            if c[1][:1] in ("A", "C") and status_bucket(c[5]) == "⏳ 준비·대기·계획":
                L.append(f"| {c[0]} | {c[1]} | {c[2][:45]} | {c[4][:30]} |")
        L.append("")
    return "\n".join(L)


# ── git ──────────────────────────────────────────────────────────────────
def git(*args, env=None, input=None):
    e = dict(os.environ)
    if env:
        e.update(env)
    r = subprocess.run(["git", *args], env=e, input=input, capture_output=True, text=True)
    if r.returncode:
        sys.exit(f"git {' '.join(args)} 실패: {r.stderr.strip()}")
    return r.stdout.strip()


def dated_env(when):
    stamp = when.strftime("%Y-%m-%dT%H:%M:00") + KST
    return {
        "GIT_AUTHOR_NAME": AUTHOR[0], "GIT_AUTHOR_EMAIL": AUTHOR[1], "GIT_AUTHOR_DATE": stamp,
        "GIT_COMMITTER_NAME": COMMITTER[0], "GIT_COMMITTER_EMAIL": COMMITTER[1], "GIT_COMMITTER_DATE": stamp,
    }


def day_message(kind, day, added, changed, areas):
    parts = []
    for tab in TABS:
        a, c = added.get(tab, 0), changed.get(tab, 0)
        if a or c:
            parts.append(f"{tab} " + " · ".join(x for x in (f"+{a}" if a else "", f"~{c}" if c else "") if x))
    head = f"로드맵 {kind} {day}: " + " / ".join(parts)
    top = ", ".join(f"{k} {v}" for k, v in areas.most_common(4))
    body = (
        f"\n\n영역: {top}\n\n"
        "시트 작성·수정일 기준 소급. 행은 그 날짜의 마지막 수정 내용으로 들어간다 —\n"
        "시트는 행별 수정 이력을 남기지 않으므로 중간 판본은 복원할 수 없다."
    )
    return head + body + TRAILER


def cmd_export(a):
    data = load(a.xlsx, not a.no_redact)
    os.makedirs(OUT_DIR, exist_ok=True)
    for tab, (header, rows) in data.items():
        with open(os.path.join(OUT_DIR, TABS[tab][0]), "w", encoding="utf-8") as f:
            f.write(render(header, rows))
        print(f"{tab}: {len(rows)}행")
    write_summary(load_tsv_all())


def cmd_backfill(a):
    data = load(a.xlsx, not a.no_redact)
    by_day = collections.defaultdict(list)  # day -> [(tab, row)]
    for tab, (_, rows) in data.items():
        for r in rows:
            by_day[r[2].date()].append((tab, r))
    state = {tab: [] for tab in data}
    parent = None
    with tempfile.TemporaryDirectory() as td:
        idx = {"GIT_INDEX_FILE": os.path.join(td, "index")}
        for day in sorted(by_day):
            added, areas = collections.Counter(), collections.Counter()
            latest = max(r[2] for _, r in by_day[day])
            for tab, r in by_day[day]:
                state[tab].append(r)
                added[tab] += 1
                areas[r[1][1]] += 1
            for tab, rows in state.items():
                if not rows:
                    continue
                rows.sort(key=lambda r: (int(r[0]) if r[0].isdigit() else 10**9, r[0]))
                blob = git("hash-object", "-w", "--stdin", input=render(data[tab][0], rows))
                git("update-index", "--add", "--cacheinfo", f"100644,{blob},{OUT_DIR}/{TABS[tab][0]}", env=idx)
            tree = git("write-tree", env=idx)
            args = ["commit-tree", tree, "-m", day_message("소급", day, added, {}, areas)]
            if parent:
                args[2:2] = ["-p", parent]
            parent = git(*args, env=dated_env(latest))
            print(f"{day}  {sum(added.values()):>3}행  {parent[:10]}")
    print(parent)


def cmd_sync(a):
    data = load(a.xlsx, not a.no_redact)
    if git("status", "--porcelain", "--", OUT_DIR):
        sys.exit(f"{OUT_DIR} 에 커밋 안 된 변경이 있다. 먼저 정리한다.")
    plan = collections.defaultdict(list)  # day -> [(tab, row, is_new)]
    removed = {}
    current = {}
    for tab, (header, rows) in data.items():
        path = os.path.join(OUT_DIR, TABS[tab][0])
        _, old = read_tsv(path)
        current[tab] = dict(old)
        new_keys = set()
        for r in rows:
            new_keys.add(r[0])
            if old.get(r[0]) != r[1]:
                plan[r[2].date()].append((tab, r, r[0] not in old))
        removed[tab] = sorted(set(old) - new_keys, key=lambda k: int(k) if k.isdigit() else 10**9)
    if not plan and not any(removed.values()):
        print("바뀐 행 없음")
        return
    for day in sorted(plan):
        n = collections.Counter(("신규" if new else "수정") for _, _, new in plan[day])
        print(f"{day}  신규 {n['신규']} · 수정 {n['수정']}")
    for tab, ks in removed.items():
        if ks:
            print(f"{tab}: 시트에서 사라진 행 {len(ks)} — {', '.join(ks[:10])}")
    if a.dry_run:
        return

    def write(tab):
        header = data[tab][0]
        rows = sorted(current[tab].items(), key=lambda kv: int(kv[0]) if kv[0].isdigit() else 10**9)
        with open(os.path.join(OUT_DIR, TABS[tab][0]), "w", encoding="utf-8") as f:
            f.write("\n".join(["\t".join(header)] + ["\t".join(v) for _, v in rows]) + "\n")

    os.makedirs(OUT_DIR, exist_ok=True)
    for day in sorted(plan):
        added, changed, areas = collections.Counter(), collections.Counter(), collections.Counter()
        latest = max(r[2] for _, r, _ in plan[day])
        for tab, r, new in plan[day]:
            current[tab][r[0]] = r[1]
            (added if new else changed)[tab] += 1
            areas[r[1][1]] += 1
        for tab in {t for t, _, _ in plan[day]}:
            write(tab)
        git("add", OUT_DIR)
        git("commit", "-q", "-m", day_message("동기화", day, added, changed, areas), env=dated_env(latest))
    if any(removed.values()):
        for tab, ks in removed.items():
            for k in ks:
                current[tab].pop(k, None)
            write(tab)
        git("add", OUT_DIR)
        msg = "로드맵 동기화: 시트에서 사라진 행 반영\n\n" + "\n".join(
            f"{t}: {', '.join(ks)}" for t, ks in removed.items() if ks) + TRAILER
        git("commit", "-q", "-m", msg)
    write_summary(load_tsv_all())
    if git("status", "--porcelain", "--", OUT_DIR):
        git("add", OUT_DIR)
        git("commit", "-q", "-m", "로드맵 요약 갱신" + TRAILER)
    print("완료 — git log --date=iso -- orangpro/roadmap 로 확인")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("add", help="새 행 기입(평소 쓰는 것)")
    s.add_argument("tab", help="로드맵 | 로드맵(미래) | 100억")
    s.add_argument("delta", help="헤더 없는 TSV. 열 수는 그 탭과 같다")
    s.add_argument("--date", help="날짜 열이 없거나 빈 행의 날짜 YYYY-MM-DD (기본 오늘)")
    s.add_argument("--why", help="커밋 본문에 붙일 한 줄 사유")
    s.add_argument("--loose", action="store_true", help="번호 연속 검사 끔")
    s.add_argument("--dry-run", action="store_true")
    s.add_argument("--no-redact", action="store_true", help="마스킹 끔 — 비공개 저장소에서만")
    s.set_defaults(fn=cmd_add)
    s = sub.add_parser("check", help="TSV 검사")
    s.add_argument("--loose", action="store_true", help="번호 연속 검사 끔")
    s.set_defaults(fn=cmd_check)
    s = sub.add_parser("summary", help="SUMMARY.md 재생성")
    s.set_defaults(fn=cmd_summary)
    for name, fn in (("export", cmd_export), ("backfill", cmd_backfill), ("sync", cmd_sync)):
        s = sub.add_parser(name, help="시트(XLSX)에서 가져오기")
        s.add_argument("xlsx")
        s.add_argument("--no-redact", action="store_true", help="마스킹 끔 — 비공개 저장소에서만")
        if name == "sync":
            s.add_argument("--dry-run", action="store_true")
        s.set_defaults(fn=fn)
    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
